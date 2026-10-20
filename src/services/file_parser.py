import io
import base64
from pathlib import Path
import structlog

logger = structlog.get_logger(__name__)

# Спроба імпортувати pypdf або PyPDF2
try:
    import pypdf
    _pdf_reader_cls = pypdf.PdfReader
except ImportError:
    try:
        import PyPDF2 as pypdf
        _pdf_reader_cls = pypdf.PdfReader
    except ImportError:
        _pdf_reader_cls = None

try:
    from docx import Document as DocxDocument
except ImportError:
    DocxDocument = None


class OCRUnavailableError(Exception):
    """OCR не вдалося виконати (ліміт Groq, мережа). Це не означає, що документ порожній."""


def _ocr_image_with_groq(image_bytes: bytes, filename: str = "", strict: bool = False) -> str:
    """
    Виконує оптичне розпізнавання тексту (OCR) на зображенні через мультимодальну модель Groq (qwen/qwen3.8-27b).
    """
    log = logger.bind(filename=filename, image_size=len(image_bytes))
    log.info("groq_pdf_ocr_started")

    try:
        from groq import Groq
        from src.core.config import settings

        client = Groq(
            api_key=settings.groq_api_key.get_secret_value()
            if hasattr(settings.groq_api_key, "get_secret_value")
            else settings.groq_api_key
        )

        img_b64 = base64.b64encode(image_bytes).decode("utf-8")

        # Визначаємо MIME-тип
        mime_type = "image/jpeg"
        if image_bytes.startswith(b"\x89PNG"):
            mime_type = "image/png"

        prompt = (
            "Твоє завдання — переписати весь текст із цього документа українською мовою точно як на зображенні, "
            "слово в слово, зберігаючи всі назви, цифри та поради.\n"
            "Відповідай ВИКЛЮЧНО чистим текстом самого документа без жодних вступних фраз (не пиши 'Ось текст:', "
            "'Текст з документа:' тощо) та без власних коментарів."
        )

        response = client.chat.completions.create(
            model="qwen/qwen3.8-27b",
            messages=[{
                "role": "user",
                "content": [
                    {"type": "text", "text": prompt},
                    {"type": "image_url", "image_url": {"url": f"data:{mime_type};base64,{img_b64}"}}
                ]
            }],
            temperature=0.1
        )

        text = response.choices[0].message.content.strip()
        log.info("groq_pdf_ocr_completed", extracted_chars=len(text))
        return text

    except Exception as e:
        log.error("groq_pdf_ocr_failed", error=str(e))
        if strict:
            raise OCRUnavailableError(str(e)) from e
        return ""


def _extract_images_from_pdf_page(page) -> list[bytes]:
    """Витягує сирі байти зображень зі сторінки PDF."""
    images = []
    try:
        resources = page.get("/Resources", {})
        if hasattr(resources, "get_object"):
            resources = resources.get_object()

        xobjects = resources.get("/XObject", {})
        if hasattr(xobjects, "get_object"):
            xobjects = xobjects.get_object()

        if isinstance(xobjects, dict):
            for name, obj in xobjects.items():
                if hasattr(obj, "get_object"):
                    obj = obj.get_object()
                if isinstance(obj, dict) and obj.get("/Subtype") == "/Image":
                    data = obj.get_data()
                    if data:
                        images.append(data)
    except Exception as e:
        logger.warning("extract_images_from_pdf_page_failed", error=str(e))
    return images


def extract_text(file_bytes: bytes, filename: str, strict_ocr: bool = False) -> str:
    """
    Перетворює сирі байти завантаженого файлу на чистий текст.
    Підтримувані формати: .txt, .pdf, .docx.
    Якщо PDF є сканованим зображенням без текстового шару — автоматично викликає Groq OCR.
    """
    ext = Path(filename).suffix.lower()
    log = logger.bind(filename=filename, extension=ext, byte_size=len(file_bytes))
    log.info("extracting_text_started")

    text = ""

    # 1. Текстові файли (.txt)
    if ext == ".txt":
        try:
            text = file_bytes.decode("utf-8")
            log.debug("text_decoded_utf8")
        except UnicodeDecodeError:
            try:
                text = file_bytes.decode("cp1251")
                log.debug("text_decoded_cp1251")
            except UnicodeDecodeError as e:
                log.warning("text_decode_failed_fallback_ignore", error=str(e))
                text = file_bytes.decode("utf-8", errors="ignore")

    # 2. PDF файли (.pdf)
    elif ext == ".pdf":
        if _pdf_reader_cls is None:
            log.error("pdf_library_not_installed")
            raise RuntimeError("Бібліотека pypdf не встановлена")

        try:
            pdf_stream = io.BytesIO(file_bytes)
            reader = _pdf_reader_cls(pdf_stream)
            pages_text = []

            for idx, page in enumerate(reader.pages):
                page_content = page.extract_text()
                if page_content and page_content.strip():
                    pages_text.append(page_content.strip())

            text = "\n".join(pages_text).strip()

            # Якщо текстовий шар порожній (сканований PDF-документ) — витягуємо зображення та робимо OCR
            if not text:
                log.info("pdf_has_no_text_layer_attempting_ocr", total_pages=len(reader.pages))
                ocr_pages = []
                for idx, page in enumerate(reader.pages):
                    page_images = _extract_images_from_pdf_page(page)
                    for img_bytes in page_images:
                        ocr_result = _ocr_image_with_groq(img_bytes, filename=filename, strict=strict_ocr)
                        if ocr_result:
                            ocr_pages.append(ocr_result)

                text = "\n\n".join(ocr_pages).strip()
                if text:
                    log.info("pdf_extracted_via_ocr_successfully", total_pages=len(reader.pages), extracted_chars=len(text))
                else:
                    log.warning("pdf_ocr_returned_no_text")
            else:
                log.info("pdf_extracted_successfully", total_pages=len(reader.pages), extracted_chars=len(text))

        except Exception as e:
            log.error("pdf_extraction_error", error=str(e))
            raise

    # 3. Word файли (.docx)
    elif ext in [".docx", ".doc"]:
        if DocxDocument is None:
            log.error("docx_library_not_installed")
            raise RuntimeError("Бібліотека python-docx не встановлена")

        try:
            docx_stream = io.BytesIO(file_bytes)
            doc = DocxDocument(docx_stream)
            paragraphs_text = [p.text for p in doc.paragraphs if p.text]
            text = "\n".join(paragraphs_text).strip()
            log.info("docx_extracted_successfully", total_paragraphs=len(doc.paragraphs), extracted_chars=len(text))
        except Exception as e:
            log.error("docx_extraction_error", error=str(e))
            raise

    else:
        log.warning("unsupported_file_extension")
        raise ValueError(f"Непідтримуваний формат файлу: {ext}")

    clean_text = text.strip()
    log.info("extracting_text_completed", extracted_chars=len(clean_text))
    return clean_text